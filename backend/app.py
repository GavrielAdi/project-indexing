# ======================================================================
# BAGIAN 0: IMPOR DAN KONFIGURASI
# ======================================================================
import os
import re
import io
import json
import tempfile
from datetime import datetime
from bson import ObjectId
from bson.errors import InvalidId

from flask import Flask, request, jsonify
from flask_cors import CORS
from werkzeug.utils import secure_filename
from dotenv import load_dotenv
from markupsafe import escape
# Impor Pustaka Anda
from docx import Document
import fitz
from pdf2image import convert_from_path
import pytesseract
import openpyxl
import magic

from pymongo import MongoClient
import gridfs

# --- KONFIGURASI APLIKASI ---
app = Flask(__name__)
CORS(app, resources={r"/*": {"origins": "*"}})
load_dotenv()

ALLOWED_EXTENSIONS = {'docx', 'pdf', 'xlsx', 'txt'}
ALLOWED_MIME_TYPES = {
    'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    'pdf': 'application/pdf',
    'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'txt': 'text/plain'
}


# --- KONEKSI KE MONGODB ---
MONGO_URI = os.getenv("MONGO_URI")
try:
    print("--- Mencoba koneksi ke MongoDB... ---")
    if not MONGO_URI:
        raise ValueError("MONGO_URI tidak ditemukan di environment variables. Buat file .env")
    client = MongoClient(MONGO_URI, serverSelectionTimeoutMS=5000)
    client.server_info()
    db = client['dokumen_cerdas_db']
    fs = gridfs.GridFS(db)
    file_metadata_collection = db['file_metadata']
    print("--- Koneksi ke MongoDB berhasil ---")
except Exception as e:
    print(f"--- KRITIS: GAGAL KONEKSI KE MONGODB. APLIKASI BERHENTI. ---")
    print(f"Error: {e}")
    exit()


# ======================================================================
# BAGIAN 1: KELAS-KELAS ALGORITMA INDEXING
# ======================================================================
class PositionalIndex:
    def __init__(self): self.indeks, self.dokumen = {}, []
    def build(self, tokens):
        self.dokumen = tokens
        for i, token in enumerate(tokens):
            if token not in self.indeks: self.indeks[token] = []
            self.indeks[token].append(i)
    def search_phrase(self, frasa):
        words = frasa.strip().lower().split()
        if not words or words[0] not in self.indeks: return []
        posisi_awal = self.indeks.get(words[0], [])
        hasil_posisi = []
        for pos in posisi_awal:
            cocok = True
            for i in range(1, len(words)):
                if (words[i] not in self.indeks) or ((pos + i) not in self.indeks[words[i]]):
                    cocok = False; break
            if cocok: hasil_posisi.append(pos)
        return hasil_posisi
    
    def to_dict(self):
        return {'indeks': self.indeks, 'dokumen': self.dokumen}

    @classmethod
    def from_dict(cls, data):
        instance = cls()
        instance.indeks = data.get('indeks', {})
        instance.dokumen = data.get('dokumen', [])
        return instance

class TrieNode:
    def __init__(self): self.children, self.is_end_of_word = {}, False

class Trie:
    def __init__(self): self.root = TrieNode()
    def insert(self, word):
        node = self.root
        for char in word:
            if char not in node.children: node.children[char] = TrieNode()
            node = node.children[char]
        node.is_end_of_word = True
    def build(self, tokens):
        for token in set(tokens): self.insert(token)
    def _dfs(self, node, prefix):
        hasil = []
        if node.is_end_of_word: hasil.append(prefix)
        for char, next_node in node.children.items():
            hasil.extend(self._dfs(next_node, prefix + char))
        return hasil
    def get_words_with_prefix(self, prefix):
        node = self.root
        for char in prefix:
            if char not in node.children: return []
            node = node.children[char]
        return self._dfs(node, prefix)


# ======================================================================
# BAGIAN 2: FUNGSI-FUNGSI PEMBANTU
# ======================================================================
def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def process_file_from_bytes(file_bytes, filename):
    extension = os.path.splitext(filename)[1].lower()
    full_text = ""
    file_stream = io.BytesIO(file_bytes)
    try:
        if extension == '.docx':
            doc = Document(file_stream)
            full_text = " ".join([para.text for para in doc.paragraphs])
        elif extension == '.pdf':
            try:
                with fitz.open(stream=file_bytes, filetype="pdf") as doc:
                    text = "".join(page.get_text() for page in doc)
                if len(text.strip()) > 50: full_text = text
            except Exception: full_text = ""
            if not full_text:
                print("Info: PDF butuh OCR. Membuat file temporer...")
                with tempfile.NamedTemporaryFile(suffix=".pdf", delete=True) as tmp:
                    tmp.write(file_bytes)
                    tmp.flush()
                    images = convert_from_path(tmp.name)
                    ocr_text = "".join([pytesseract.image_to_string(img, lang='ind+eng') for img in images])
                    full_text = ocr_text
        elif extension == '.xlsx':
            workbook = openpyxl.load_workbook(file_stream)
            text_parts = [str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row if cell.value]
            full_text = " ".join(text_parts)
        elif extension == '.txt':
            full_text = file_stream.read().decode('utf-8', errors='ignore')
        if full_text:
            cleaned_text = re.sub(r'[^\w\s]', '', full_text.lower())
            return cleaned_text.split()
    except Exception as e:
        app.logger.error(f"Error saat memproses {filename}: {e}")
    return None

# ======================================================================
# BAGIAN 3: INISIALISASI STATE DI MEMORI
# ======================================================================
pos_index = PositionalIndex()
trie_index = Trie()
current_indexed_file = "Tidak ada dokumen yang dipilih."

# ======================================================================
# BAGIAN 4: API ENDPOINTS
# ======================================================================

@app.route('/status', methods=['GET'])
def status():
    return jsonify({'indexed_file': current_indexed_file})

@app.route('/upload', methods=['POST'])
def upload_file_to_db():
    if 'file' not in request.files:
        return jsonify({'error': 'No file part'}), 400
    file = request.files['file']
    uploaded_by = request.form.get('uploaded_by', 'Anonim')
    tags_string = request.form.get('tags', '')
    if file.filename == '' or not allowed_file(file.filename):
        return jsonify({'error': 'File tidak dipilih atau tipe tidak diizinkan'}), 400
    filename = secure_filename(file.filename)
    if file_metadata_collection.find_one({'filename': filename}):
        return jsonify({'error': f'File dengan nama {filename} sudah ada di database.'}), 409
    file_content = file.read()
    detected_mime_type = magic.from_buffer(file_content, mime=True)
    file_ext = filename.rsplit('.', 1)[1].lower()
    expected_mime_type = ALLOWED_MIME_TYPES.get(file_ext)
    if detected_mime_type != expected_mime_type:
        return jsonify({'error': f'Konten file tidak cocok dengan ekstensi .{file_ext}. Terdeteksi: {detected_mime_type}'}), 400
    try:
        gridfs_id = fs.put(file_content, filename=filename)
    except Exception as e:
        app.logger.error(f"GridFS put error: {e}")
        return jsonify({'error': 'Gagal menyimpan file ke database.'}), 500
    tokens = process_file_from_bytes(file_content, filename)
    if not tokens:
        fs.delete(gridfs_id)
        return jsonify({'error': 'Gagal memproses teks dari file'}), 500
    pos_index_obj = PositionalIndex()
    pos_index_obj.build(tokens)
    index_data_dict = pos_index_obj.to_dict()
    tags_list = [tag.strip() for tag in tags_string.split(',') if tag.strip()]
    now = datetime.now()
    file_metadata_collection.insert_one({
        'gridfs_id': gridfs_id,
        'filename': filename,
        'upload_date': now,
        'last_modified_date': now,
        'index_data': index_data_dict,
        'uploaded_by': uploaded_by,
        'tags': tags_list
    })
    return jsonify({'message': f'File {filename} berhasil disimpan ke database dan di-indeks.'}), 201

@app.route('/documents', methods=['GET'])
def get_documents_from_db():
    try:
        documents = file_metadata_collection.find({}, {"filename": 1, "upload_date": 1, "uploaded_by": 1, "tags": 1, "last_modified_date": 1}).sort("last_modified_date", -1)
        doc_list = []
        for doc in documents:
            doc_list.append({
                'id': str(doc['_id']), 
                'filename': doc['filename'],
                'upload_date': doc['upload_date'].strftime("%Y-%m-%d %H:%M:%S"),
                'last_modified_date': doc.get('last_modified_date', doc['upload_date']).strftime("%Y-%m-%d %H:%M:%S"),
                'uploaded_by': doc.get('uploaded_by', 'N/A'),
                'tags': doc.get('tags', [])
            })
        return jsonify(doc_list)
    except Exception as e:
        app.logger.error(f"Get documents error: {e}")
        return jsonify({'error': 'Gagal mengambil daftar dokumen.'}), 500

@app.route('/tags', methods=['GET'])
def get_all_tags():
    try:
        all_tags = file_metadata_collection.distinct('tags')
        all_tags.sort()
        return jsonify(all_tags)
    except Exception as e:
        app.logger.error(f"Get tags error: {e}")
        return jsonify({'error': 'Gagal mengambil daftar tag.'}), 500

@app.route('/switch_document/<doc_id>', methods=['POST'])
def switch_document(doc_id):
    global pos_index, trie_index, current_indexed_file
    try:
        obj_id = ObjectId(doc_id)
        doc_metadata = file_metadata_collection.find_one({'_id': obj_id})
        if not doc_metadata:
            return jsonify({'error': 'Dokumen tidak ditemukan'}), 404
        
        index_data = doc_metadata['index_data']
        pos_index = PositionalIndex.from_dict(index_data)
        trie_index = Trie()
        trie_index.build(set(pos_index.dokumen))
        
        current_indexed_file = doc_metadata['filename']
        
        print(f"--- Berpindah ke dokumen aktif: {current_indexed_file} ---")
        return jsonify({'message': f'Berhasil beralih ke {current_indexed_file}', 'filename': current_indexed_file})
    except (InvalidId, TypeError):
        return jsonify({'error': 'Format ID dokumen tidak valid.'}), 400
    except Exception as e:
        app.logger.error(f"Switch document error: {e}")
        return jsonify({'error': 'Gagal beralih dokumen.'}), 500

@app.route('/search', methods=['GET'])
def search():
    query = request.args.get('query', '')
    if not query or current_indexed_file == "Tidak ada dokumen yang dipilih.":
        return jsonify({'count': 0, 'snippets': []})
    
    positions = pos_index.search_phrase(query)
    snippets = []
    if positions:
        query_len = len(query.strip().lower().split())
        for pos in positions[:5]:
            start = max(0, pos - 15)
            end = min(len(pos_index.dokumen), pos + query_len + 15)
            snippet_html = " ".join(pos_index.dokumen[start:end])
            
            safe_query_words = [escape(word) for word in query.strip().lower().split()]
            for word in safe_query_words:
                pattern = re.compile(r'\b(' + re.escape(word) + r')\b', re.IGNORECASE)
                snippet_html = pattern.sub(r'<strong>\1</strong>', snippet_html)
            snippets.append(f"...{snippet_html}...")
    return jsonify({'count': len(positions), 'snippets': snippets})

@app.route('/autocomplete', methods=['GET'])
def autocomplete():
    prefix = request.args.get('prefix', '')
    if not prefix or current_indexed_file == "Tidak ada dokumen yang dipilih.":
        return jsonify([])
    return jsonify(trie_index.get_words_with_prefix(prefix))

@app.route('/document/<doc_id>', methods=['DELETE'])
def delete_document(doc_id):
    global current_indexed_file, pos_index, trie_index
    try:
        obj_id = ObjectId(doc_id)
        doc_metadata = file_metadata_collection.find_one({'_id': obj_id})
        if not doc_metadata:
            return jsonify({'error': 'Dokumen tidak ditemukan'}), 404

        gridfs_id = doc_metadata['gridfs_id']
        filename_to_delete = doc_metadata['filename']

        fs.delete(gridfs_id)
        file_metadata_collection.delete_one({'_id': obj_id})

        if filename_to_delete == current_indexed_file:
            current_indexed_file = "Tidak ada dokumen yang dipilih."
            pos_index = PositionalIndex()
            trie_index = Trie()

        return jsonify({'message': f'Dokumen {filename_to_delete} berhasil dihapus.'}), 200
    except (InvalidId, TypeError):
        return jsonify({'error': 'Format ID dokumen tidak valid.'}), 400
    except Exception as e:
        app.logger.error(f"Delete document error: {e}")
        return jsonify({'error': 'Gagal menghapus dokumen.'}), 500

@app.route('/document/<doc_id>', methods=['PUT'])
def update_document(doc_id):
    try:
        obj_id = ObjectId(doc_id)
    except (InvalidId, TypeError):
        return jsonify({'error': 'Format ID dokumen tidak valid.'}), 400

    try:
        data = request.get_json()
        if not data:
            return jsonify({'error': 'Request body tidak boleh kosong.'}), 400

        updated_by = data.get('uploaded_by')
        tags_string = data.get('tags', '')

        if not updated_by or not updated_by.strip():
            return jsonify({'error': 'Nama pengunggah tidak boleh kosong.'}), 400

        tags_list = [tag.strip() for tag in tags_string.split(',') if tag.strip()]

        update_result = file_metadata_collection.update_one(
            {'_id': obj_id},
            {'$set': {
                'uploaded_by': updated_by.strip(),
                'tags': tags_list,
                'last_modified_date': datetime.now()
            }}
        )

        if update_result.matched_count == 0:
            return jsonify({'error': 'Dokumen tidak ditemukan.'}), 404

        return jsonify({'message': 'Dokumen berhasil diperbarui.'}), 200

    except Exception as e:
        app.logger.error(f"Update document error: {e}")
        return jsonify({'error': 'Gagal memperbarui dokumen.'}), 500

@app.route('/stats', methods=['GET'])
def get_stats():
    try:
        total_documents = file_metadata_collection.count_documents({})
        total_tags = len(file_metadata_collection.distinct('tags'))
        
        latest_doc = file_metadata_collection.find_one(
            {}, 
            sort=[("upload_date", -1)]
        )
        
        latest_upload = {
            'filename': latest_doc['filename'],
            'upload_date': latest_doc['upload_date'].strftime("%Y-%m-%d %H:%M:%S")
        } if latest_doc else None

        return jsonify({
            'total_documents': total_documents,
            'total_tags': total_tags,
            'latest_upload': latest_upload
        })
    except Exception as e:
        app.logger.error(f"Get stats error: {e}")
        return jsonify({'error': 'Gagal mengambil statistik.'}), 500

@app.errorhandler(500)
def internal_server_error(e):
    app.logger.error(f"Internal Server Error: {e}", exc_info=True)
    return jsonify(error="Terjadi kesalahan internal pada server. Silakan cek log."), 500

# ======================================================================
# BAGIAN 4: MENJALANKAN APLIKASI
# ======================================================================
if __name__ == '__main__':
    print("--- Server dimulai, siap menerima koneksi ---")
    app.run(host='0.0.0.0', port=5000, debug=True)
